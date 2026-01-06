"""Reporting and export functionality.

This module provides functions to generate reports and export data
in various formats including CSV, Excel, and JSON.
"""

import csv
import json
from datetime import datetime
from decimal import Decimal
from io import BytesIO, StringIO
from pathlib import Path
from typing import Any, Optional, Union

import structlog
from sqlalchemy import text
from sqlalchemy.orm import Session

logger = structlog.get_logger()


class DecimalEncoder(json.JSONEncoder):
    """JSON encoder that handles Decimal types."""

    def default(self, obj):
        if isinstance(obj, Decimal):
            return float(obj)
        return super().default(obj)


class ReportGenerator:
    """Generate reports and export data.

    Provides functionality to create market summary reports, valuation
    distributions, and physician cohort analysis with export to
    various formats.

    Example:
        >>> generator = ReportGenerator()
        >>> csv_data = generator.export_valuations_csv(session)
        >>> with open("valuations.csv", "w") as f:
        ...     f.write(csv_data)
    """

    def generate_market_summary(
        self,
        session: Session,
        states: Optional[list[str]] = None,
    ) -> dict:
        """Generate market summary report.

        Args:
            session: SQLAlchemy session.
            states: Filter by states (optional).

        Returns:
            Dictionary with market summary data.
        """
        report = {
            "generated_at": datetime.utcnow().isoformat(),
            "report_type": "market_summary",
        }

        # Overall metrics
        overall_query = """
            SELECT
                COUNT(DISTINCT p.id) as practice_count,
                COUNT(DISTINCT ph.npi) as physician_count,
                SUM(v.valuation_amount) as total_valuation,
                AVG(v.valuation_amount) as avg_valuation,
                AVG(p.payer_mix_commercial) as avg_commercial_payer
            FROM practices p
            LEFT JOIN valuations v ON p.id = v.practice_id
            LEFT JOIN practice_physicians pp ON p.id = pp.practice_id
            LEFT JOIN physicians ph ON pp.physician_npi = ph.npi
        """

        params = {}
        if states:
            placeholders = ", ".join(f":state_{i}" for i in range(len(states)))
            overall_query += f" WHERE p.state IN ({placeholders})"
            for i, state in enumerate(states):
                params[f"state_{i}"] = state

        result = session.execute(text(overall_query), params)
        row = result.fetchone()

        report["overview"] = {
            "total_practices": row[0] or 0,
            "total_physicians": row[1] or 0,
            "total_market_value": float(row[2] or 0),
            "avg_practice_value": float(row[3] or 0),
            "avg_commercial_payer_pct": float(row[4] or 0),
        }

        # Breakdown by state
        state_query = text("""
            SELECT
                p.state,
                COUNT(DISTINCT p.id) as practice_count,
                SUM(v.valuation_amount) as total_valuation,
                AVG(v.valuation_amount) as avg_valuation
            FROM practices p
            LEFT JOIN valuations v ON p.id = v.practice_id
            WHERE p.state IS NOT NULL
            GROUP BY p.state
            ORDER BY total_valuation DESC
        """)

        state_result = session.execute(state_query)
        report["by_state"] = [
            {
                "state": row[0],
                "practice_count": row[1],
                "total_valuation": float(row[2] or 0),
                "avg_valuation": float(row[3] or 0),
            }
            for row in state_result.fetchall()
        ]

        # Breakdown by specialty
        specialty_query = text("""
            SELECT
                p.specialty_primary,
                COUNT(DISTINCT p.id) as practice_count,
                AVG(v.valuation_amount) as avg_valuation,
                AVG(v.revenue_multiple) as avg_multiple
            FROM practices p
            LEFT JOIN valuations v ON p.id = v.practice_id
            WHERE p.specialty_primary IS NOT NULL
            GROUP BY p.specialty_primary
            ORDER BY avg_valuation DESC
        """)

        specialty_result = session.execute(specialty_query)
        report["by_specialty"] = [
            {
                "specialty": row[0],
                "practice_count": row[1],
                "avg_valuation": float(row[2] or 0),
                "avg_revenue_multiple": float(row[3] or 0),
            }
            for row in specialty_result.fetchall()
        ]

        return report

    def generate_valuation_distribution(
        self,
        session: Session,
        bin_size: int = 500000,
    ) -> dict:
        """Generate valuation distribution analysis.

        Args:
            session: SQLAlchemy session.
            bin_size: Size of valuation bins (default $500K).

        Returns:
            Dictionary with distribution data.
        """
        query = text(f"""
            SELECT
                FLOOR(valuation_amount / {bin_size}) * {bin_size} as bin_start,
                COUNT(*) as count,
                SUM(valuation_amount) as total_value
            FROM valuations
            GROUP BY bin_start
            ORDER BY bin_start
        """)

        result = session.execute(query)
        rows = result.fetchall()

        bins = []
        for row in rows:
            bin_start = float(row[0] or 0)
            bins.append({
                "range": f"${bin_start/1000000:.1f}M - ${(bin_start + bin_size)/1000000:.1f}M",
                "min_value": bin_start,
                "max_value": bin_start + bin_size,
                "count": row[1],
                "total_value": float(row[2] or 0),
            })

        return {
            "generated_at": datetime.utcnow().isoformat(),
            "report_type": "valuation_distribution",
            "bin_size": bin_size,
            "distribution": bins,
            "total_practices": sum(b["count"] for b in bins),
            "total_value": sum(b["total_value"] for b in bins),
        }

    def generate_cohort_analysis(
        self,
        session: Session,
    ) -> dict:
        """Generate physician cohort analysis.

        Args:
            session: SQLAlchemy session.

        Returns:
            Dictionary with cohort analysis data.
        """
        # Segment distribution
        segment_query = text("""
            SELECT
                segment_type,
                COUNT(*) as count,
                AVG(acquisition_likelihood) as avg_likelihood
            FROM physician_segments
            GROUP BY segment_type
            ORDER BY avg_likelihood DESC
        """)

        segment_result = session.execute(segment_query)

        segments = [
            {
                "segment": row[0],
                "count": row[1],
                "avg_acquisition_likelihood": float(row[2] or 0),
            }
            for row in segment_result.fetchall()
        ]

        # Age distribution
        age_query = text("""
            SELECT
                CASE
                    WHEN (EXTRACT(YEAR FROM NOW()) - birth_year) < 35 THEN 'Under 35'
                    WHEN (EXTRACT(YEAR FROM NOW()) - birth_year) < 45 THEN '35-44'
                    WHEN (EXTRACT(YEAR FROM NOW()) - birth_year) < 55 THEN '45-54'
                    WHEN (EXTRACT(YEAR FROM NOW()) - birth_year) < 65 THEN '55-64'
                    ELSE '65+'
                END as age_group,
                COUNT(*) as count
            FROM physicians
            WHERE birth_year IS NOT NULL
            GROUP BY age_group
            ORDER BY age_group
        """)

        age_result = session.execute(age_query)

        age_groups = [
            {"age_group": row[0], "count": row[1]}
            for row in age_result.fetchall()
        ]

        # Experience distribution
        experience_query = text("""
            SELECT
                CASE
                    WHEN (EXTRACT(YEAR FROM NOW()) - graduation_year) <= 5 THEN '0-5 years'
                    WHEN (EXTRACT(YEAR FROM NOW()) - graduation_year) <= 10 THEN '6-10 years'
                    WHEN (EXTRACT(YEAR FROM NOW()) - graduation_year) <= 20 THEN '11-20 years'
                    WHEN (EXTRACT(YEAR FROM NOW()) - graduation_year) <= 30 THEN '21-30 years'
                    ELSE '30+ years'
                END as experience,
                COUNT(*) as count
            FROM physicians
            WHERE graduation_year IS NOT NULL
            GROUP BY experience
            ORDER BY experience
        """)

        exp_result = session.execute(experience_query)

        experience = [
            {"experience_range": row[0], "count": row[1]}
            for row in exp_result.fetchall()
        ]

        return {
            "generated_at": datetime.utcnow().isoformat(),
            "report_type": "cohort_analysis",
            "segments": segments,
            "age_distribution": age_groups,
            "experience_distribution": experience,
        }

    def export_valuations_csv(
        self,
        session: Session,
        limit: Optional[int] = None,
    ) -> str:
        """Export valuations to CSV format.

        Args:
            session: SQLAlchemy session.
            limit: Maximum rows to export.

        Returns:
            CSV data as string.
        """
        query = """
            SELECT
                p.id,
                p.name,
                p.practice_type,
                p.specialty_primary,
                p.city,
                p.state,
                p.zip_code,
                p.physician_count,
                p.patient_panel_size,
                p.revenue_annual,
                p.payer_mix_medicare,
                p.payer_mix_commercial,
                v.valuation_amount,
                v.valuation_low,
                v.valuation_high,
                v.revenue_multiple,
                v.confidence_level,
                v.calculation_timestamp
            FROM practices p
            JOIN valuations v ON p.id = v.practice_id
            ORDER BY v.valuation_amount DESC
        """

        if limit:
            query += f" LIMIT {limit}"

        result = session.execute(text(query))
        rows = result.fetchall()

        output = StringIO()
        writer = csv.writer(output)

        # Header
        writer.writerow([
            "practice_id", "name", "type", "specialty", "city", "state",
            "zip_code", "physician_count", "patient_panel", "annual_revenue",
            "medicare_pct", "commercial_pct", "valuation", "valuation_low",
            "valuation_high", "revenue_multiple", "confidence", "calculated_at"
        ])

        # Data
        for row in rows:
            writer.writerow(row)

        return output.getvalue()

    def export_to_excel(
        self,
        session: Session,
        filepath: Union[str, Path],
    ) -> None:
        """Export data to Excel format.

        Args:
            session: SQLAlchemy session.
            filepath: Output file path.
        """
        try:
            import openpyxl
            from openpyxl.utils.dataframe import dataframe_to_rows
        except ImportError:
            raise ImportError("openpyxl required for Excel export")

        wb = openpyxl.Workbook()

        # Valuations sheet
        ws_val = wb.active
        ws_val.title = "Valuations"

        val_query = text("""
            SELECT
                p.name, p.specialty_primary, p.state,
                v.valuation_amount, v.revenue_multiple, v.confidence_level
            FROM practices p
            JOIN valuations v ON p.id = v.practice_id
            ORDER BY v.valuation_amount DESC
            LIMIT 1000
        """)

        result = session.execute(val_query)
        ws_val.append(["Practice", "Specialty", "State", "Valuation", "Multiple", "Confidence"])
        for row in result.fetchall():
            ws_val.append(list(row))

        # Summary sheet
        ws_sum = wb.create_sheet("Summary")

        summary_query = text("""
            SELECT
                state,
                COUNT(*) as practices,
                SUM(valuation_amount) as total_value,
                AVG(valuation_amount) as avg_value
            FROM practices p
            JOIN valuations v ON p.id = v.practice_id
            GROUP BY state
            ORDER BY total_value DESC
        """)

        result = session.execute(summary_query)
        ws_sum.append(["State", "Practices", "Total Value", "Avg Value"])
        for row in result.fetchall():
            ws_sum.append(list(row))

        wb.save(filepath)
        logger.info("excel_export_complete", filepath=str(filepath))

    def export_to_json(
        self,
        session: Session,
        report_type: str = "market_summary",
    ) -> str:
        """Export report to JSON format.

        Args:
            session: SQLAlchemy session.
            report_type: Type of report to generate.

        Returns:
            JSON string.
        """
        if report_type == "market_summary":
            data = self.generate_market_summary(session)
        elif report_type == "valuation_distribution":
            data = self.generate_valuation_distribution(session)
        elif report_type == "cohort_analysis":
            data = self.generate_cohort_analysis(session)
        else:
            raise ValueError(f"Unknown report type: {report_type}")

        return json.dumps(data, cls=DecimalEncoder, indent=2)
